"""日付のテスト用ファイルを作る（実行: python3 test/make_date_fixtures.py）
必要なライブラリ: pip install openpyxl xlwt
"""
import datetime
import os

import openpyxl
import xlwt

D = os.path.join(os.path.dirname(__file__), "fixtures")

# ---- .xlsx：いろいろな表示形式の日付と数値 ----
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Sheet1"
for ref, value, fmt in [
    ("A1", datetime.date(2026, 10, 15), "mm-dd-yy"),                  # 組み込み 14（短い日付形式）
    ("A2", datetime.date(2026, 10, 15), 'yyyy"年"m"月"d"日"'),        # ユーザー定義
    ("A3", datetime.datetime(2026, 10, 15, 9, 30), "yyyy/m/d h:mm"),  # 日付と時刻
    ("A4", 1234, "#,##0"),                                            # 日付ではない数値
    ("A5", datetime.date(2026, 10, 15), '[$-411]ggge"年"m"月"d"日"'), # 和暦
    ("A6", datetime.time(9, 30), "h:mm"),                             # 時刻だけ
]:
    ws[ref] = value
    ws[ref].number_format = fmt
wb.save(os.path.join(D, "dates.xlsx"))

# ---- .xlsx：1904 年を起点とするブック ----
wb = openpyxl.Workbook()
wb.epoch = openpyxl.utils.datetime.CALENDAR_MAC_1904
wb.active["A1"] = datetime.date(2026, 10, 15)
wb.active["A1"].number_format = "yyyy/m/d"
wb.save(os.path.join(D, "dates-1904.xlsx"))

# ---- .xls：.xlsx と同じ規則で照合されるかを確かめる ----
wb = xlwt.Workbook()
ws = wb.add_sheet("Sheet1")
ws.write(0, 0, datetime.date(2026, 10, 15), xlwt.easyxf(num_format_str="M/D/YY"))  # 組み込み 14
ws.write(1, 0, 1234, xlwt.easyxf(num_format_str="#,##0"))
ws.write(2, 0, datetime.date(2026, 10, 15), xlwt.easyxf(num_format_str='yyyy"年"m"月"d"日"'))
ws.write(3, 0, datetime.date(2026, 10, 15), xlwt.easyxf(num_format_str='[$-411]ggge"年"m"月"d"日"'))
wb.save(os.path.join(D, "numbers.xls"))

print("created: dates.xlsx, dates-1904.xlsx, numbers.xls")
